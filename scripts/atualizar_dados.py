"""
Converte as tabelas da TIC Educação (Cetic.br, .xlsx) nos arquivos JSON usados pelo painel.

Uso:
    python scripts/atualizar_dados.py PASTA_GESTORES PASTA_ESCOLAS

Cada pasta deve conter os arquivos extraídos do zip do Cetic.br:
    tic_educacao_AAAA_<publico>_tabela_proporcao_vX.xlsx
    tic_educacao_AAAA_<publico>_tabela_margem_de_erro_vX.xlsx
    tic_educacao_AAAA_<publico>_tabela_total_vX.xlsx

Gera data/gestores.json, data/escolas.json e data/indicadores.json.
Requer: pip install openpyxl
"""
import glob, json, os, sys
import openpyxl

ESCOLAS = ['A1', 'A2', 'A3_1', 'A6', 'A7', 'A8', 'A19', 'A20', 'B1', 'B1B', 'B2', 'B4C', 'B14', 'B15', 'B16', 'B17', 'B18',
           'K1_1', 'C3C', 'E5A', 'E5B', 'E1D', 'E3', 'E4A', 'E4B', 'H1', 'F2', 'F3', 'F4A', 'F5', 'G4A', 'I5']
AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, '..', 'data')


def abrir(pasta, tipo):
    arqs = [f for f in glob.glob(os.path.join(pasta, '*_tabela_%s_v*.xlsx' % tipo))]
    if not arqs:
        sys.exit('Não encontrei a tabela "%s" em %s' % (tipo, pasta))
    return openpyxl.load_workbook(sorted(arqs)[-1], data_only=True)


def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else None


def ler(pasta, somente=None):
    prop, moe, tot = abrir(pasta, 'proporcao'), abrir(pasta, 'margem_de_erro'), abrir(pasta, 'total')
    tabelas = []
    for sn in prop.sheetnames:
        if somente and sn not in somente:
            continue
        ws = prop[sn]
        h3 = [ws.cell(3, c).value for c in range(3, ws.max_column + 1)]
        h4 = [ws.cell(4, c).value for c in range(3, ws.max_column + 1)]
        dois_niveis = any(h4)
        cols, grupo = [], None
        for a, b in zip(h3, h4):
            if a:
                grupo = a
            if a is None and b is None:
                cols.append(None)
            elif dois_niveis:
                cols.append((grupo, b))
            else:
                cols.append(('', a))
        idx = [i for i, c in enumerate(cols) if c]
        linhas, cat = [], None
        for r in range(5, ws.max_row + 1):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if a and str(a).startswith('Fonte'):
                break
            if a:
                cat = a
            if not cat:
                continue
            seg = b or 'Total'
            for i in idx:
                v = num(ws.cell(r, 3 + i).value)
                if v is None:
                    continue
                g, resp = cols[i]
                linhas.append({'key': '|'.join([sn, cat, seg, g or '', resp]), 'ind': sn, 'cat': cat, 'seg': seg,
                               'grupo': g or '', 'resp': resp, 'pct': v,
                               'moe': num(moe[sn].cell(r, 3 + i).value), 'n': num(tot[sn].cell(r, 3 + i).value)})
        tabelas.append({'ind': sn, 'titulo': ws.cell(1, 1).value.split(' - ', 1)[1].capitalize(), 'universo': ws.cell(2, 1).value, 'linhas': linhas})
    return tabelas


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    indicadores = []
    for publico, pasta, somente in (('gestores', sys.argv[1], None), ('escolas', sys.argv[2], ESCOLAS)):
        tabs = ler(pasta, somente)
        linhas = [l for t in tabs for l in t['linhas']]
        with open(os.path.join(SAIDA, publico + '.json'), 'w', encoding='utf-8') as f:
            json.dump(linhas, f, ensure_ascii=False, separators=(',', ':'))
        indicadores += [{'publico': publico, 'ind': t['ind'], 'titulo': t['titulo'], 'universo': t['universo']} for t in tabs]
        print('%s: %d tabelas, %d linhas' % (publico, len(tabs), len(linhas)))
    with open(os.path.join(SAIDA, 'indicadores.json'), 'w', encoding='utf-8') as f:
        json.dump(indicadores, f, ensure_ascii=False)


if __name__ == '__main__':
    main()
